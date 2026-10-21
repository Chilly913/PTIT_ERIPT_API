"""Convert the Markdown user guides to Word (.docx), one file per guide, images and tables included.

Needs python-docx (the AI service venv has it):
    PTIT_ERIPT_AI/.venv/Scripts/python PTIT_ERIPT_API/docs/huong-dan/md-to-docx.py
Output: docs/huong-dan/word/<same name>.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "word"
ACCENT = RGBColor(0x25, 0x2B, 0x6B)  # the app's primary colour family
FONT = "Calibri"

def _titles():
    out = {}
    for f in HERE.glob("*.md"):
        for ln in f.read_text(encoding="utf-8").splitlines():
            if ln.startswith("# "):
                out[f.name] = ln[2:].strip()
                break
    return out


TITLES = _titles()
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]*\)|\*[^*\s][^*]*\*)")


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), fill)
    pr.append(el)


def run_font(run, mono=False, size=None, bold=None, italic=None, color=None):
    run.font.name = "Consolas" if mono else FONT
    r = run._element.get_or_add_rPr()
    fonts = r.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), "Consolas" if mono else FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_inline(par, text, size=None, bold=False):
    """Add text with **bold**, `code`, *italic* and [links](url) to a paragraph."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run_font(par.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            run_font(r, mono=True, size=(size or 11) - 1, color=RGBColor(0x9A, 0x1F, 0x4B))
        elif part.startswith("[") and "](" in part:
            label, _, url = part[1:-1].partition("](")
            r = par.add_run(label)
            if url.startswith("http"):
                run_font(r, size=size, bold=bold, color=RGBColor(0x1D, 0x4E, 0xD8))
                r.underline = True
            else:  # links to other guides become the guide's title; folders just keep their label
                target = TITLES.get(Path(url).name)
                if target:
                    r.text = target
                run_font(r, size=size, bold=bold)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run_font(par.add_run(part[1:-1]), size=size, italic=True, bold=bold)
        else:
            run_font(par.add_run(part), size=size, bold=bold)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.15
    for name, size in (("Heading 1", 20), ("Heading 2", 15), ("Heading 3", 12.5)):
        h = doc.styles[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = ACCENT
        h.paragraph_format.space_before = Pt(16 if name != "Heading 3" else 12)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        rpr = h.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fonts.set(qn(attr), FONT)
    # page number in the footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        r._r.append(el)
    run_font(r, size=9, color=RGBColor(0x6B, 0x72, 0x80))


def add_table(doc, rows):
    n = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=n)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(n):
            cell = tbl.cell(i, j)
            cell.text = ""
            par = cell.paragraphs[0]
            par.paragraph_format.space_after = Pt(2)
            add_inline(par, row[j] if j < len(row) else "", size=10, bold=(i == 0))
            if i == 0:
                shade(cell, "E8EAF6")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc, base, alt, src):
    path = (base / src).resolve()
    if not path.exists():
        print("  missing image:", src)
        return
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(str(path), width=Cm(16))
    if alt:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_font(cap.add_run(alt), size=9, italic=True, color=RGBColor(0x6B, 0x72, 0x80))


def convert(md: Path):
    doc = Document()
    setup(doc)
    lines = md.read_text(encoding="utf-8").split("\n")
    i, first_h1 = 0, True
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):  # fenced code
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            for c in code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_after = Pt(0)
                run_font(p.add_run(c or " "), mono=True, size=9)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        m = re.match(r"^(#{1,3})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1 and first_h1:
                first_h1 = False
                t = doc.add_paragraph()
                run_font(t.add_run(text), size=26, bold=True, color=ACCENT)
                t.paragraph_format.space_after = Pt(4)
                sub = doc.add_paragraph()
                run_font(sub.add_run("Hệ thống học tập trực tuyến ERIPT LMS · Viện Kinh tế Bưu điện, Học viện Công nghệ Bưu chính Viễn thông"), size=10, italic=True, color=RGBColor(0x6B, 0x72, 0x80))
            else:
                h = doc.add_heading(level=level)
                add_inline(h, text)
            i += 1
            continue
        if re.match(r"^!\[.*\]\(.*\)\s*$", line):
            m = re.match(r"^!\[(.*)\]\((.*)\)\s*$", line)
            add_image(doc, md.parent, m.group(1), m.group(2))
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        if re.match(r"^\s*[-*]\s+", line):
            indent = len(line) - len(line.lstrip())
            p = doc.add_paragraph(style="List Bullet 2" if indent >= 2 else "List Bullet")
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, re.sub(r"^\s*[-*]\s+", "", line))
            i += 1
            continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:  # numbered list written out by hand so every list restarts at 1
            p = doc.add_paragraph()
            lvl = 1 if len(m.group(1)) >= 2 else 0
            p.paragraph_format.left_indent = Cm(0.9 + 0.8 * lvl)
            p.paragraph_format.first_line_indent = Cm(-0.6)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, f"{m.group(2)}.  {m.group(3)}")
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_inline(p, line.lstrip("> "), size=10.5)
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line.strip())
        i += 1
    OUT.mkdir(exist_ok=True)
    dest = OUT / (md.stem + ".docx")
    doc.save(dest)
    return dest, len(doc.tables), len(doc.inline_shapes)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    for md in sorted(HERE.glob("*.md")):
        dest, tables, images = convert(md)
        print(f"{dest.name}: {tables} bảng, {images} ảnh, {dest.stat().st_size // 1024} KB")
